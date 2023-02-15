import uuid
import random
from fpdf import FPDF
from docx import Document
from allure import step
import requests
from io import BytesIO
from PIL import Image
import codecs
from faker import Faker  # https://faker.readthedocs.io/en/stable/
from modules.Utils import Files
from conf_file import ALLURE_RESULTS_DIR

fake = Faker()


class Fake:

    @staticmethod
    @step('get_fake_password')
    def password():
        return fake.password()

    @staticmethod
    @step('get_fake_email')
    def email():
        return f'{Fake.string()}@gmail.com'

    @staticmethod
    @step('get_fake_email')
    def email(username=False):
        if username:
            username = username.split("@")
            return f'{username[0]}+{Fake.string()}@gmail.com'
        else:
            return f'{Fake.string()}@gmail.com'

    @staticmethod
    @step('get_fake_string')
    def string(charts: int = 10):
        letters_list = fake.random_letters(length=charts)
        string = ''.join([i for i in letters_list])
        return string

    @staticmethod
    @step('get_fake_number')
    def number(digits: int = 10):
        return fake.random_number(digits=digits)

    @staticmethod
    @step('get_fake_random_number')
    def random_number(min_num: int = 0, max_num: int = 9999):
        return fake.pyint(min_num, max_num)

    @staticmethod
    @step('random_float')
    def random_float(a: float, b: float, float_format=2):
        value = "{:." + str(float_format) + "f}"
        return value.format(random.uniform(a, b))

    @staticmethod
    @step('get_fake_sentence')
    def sentence(words: int = 10):
        return fake.sentence(nb_words=words)

    @staticmethod
    @step('get_fake_file_docx')
    def file_docx(words: int = 20):
        content = Fake.sentence(words)
        file_name = f'{uuid.uuid4()}.docx'
        path = f'{ALLURE_RESULTS_DIR}/{file_name}'
        document = Document()
        document.add_paragraph(content)
        document.save(path)
        return dict({'path': path, 'file_name': file_name, 'content': content, 'md5': Files.create_md5(path)})

    @staticmethod
    @step('get_fake_file_pdf')
    def file_pdf():
        content = 'Hello world'
        pdf = FPDF()
        pdf.add_page()
        file_name = f'{uuid.uuid4()}.pdf'
        path = f'{ALLURE_RESULTS_DIR}/{file_name}'
        pdf.set_font('Arial', 'B', 16)
        pdf.cell(40, 10, str(content))
        # noinspection PyArgumentList
        pdf.output(path, 'F')
        return dict({'path': path, 'file_name': file_name, 'content': content, 'md5': Files.create_md5(path)})

    @staticmethod
    @step('get_fake_avatar')
    def file_avatar(size: int = 100):
        file_name = f'avatar_{Fake.number()}.jpg'
        path = f'{ALLURE_RESULTS_DIR}/{file_name}'
        url = f'https://i.pravatar.cc/{size}'
        content = requests.get(url).content
        i = Image.open(BytesIO(content))
        i.save(path)
        return dict({'path': path, 'file_name': file_name, 'content': content, 'md5': Files.create_md5(path)})

    @staticmethod
    @step('file_with_size')
    def file_with_size(size_in_bytes: int):
        file_name = f'{Fake.number()}-size-{size_in_bytes}bytes.txt'
        path = f'{ALLURE_RESULTS_DIR}/{file_name}'
        f = codecs.open(path, 'w+', '1251')
        f.write('0' * size_in_bytes)
        f.close()
        return dict({'path': path, 'file_name': file_name, 'md5': Files.create_md5(path)})
